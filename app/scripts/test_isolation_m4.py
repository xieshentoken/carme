"""M4 contract tests. All inputs, names, secrets and identities are synthetic."""
from __future__ import annotations
import asyncio, copy, hashlib, io, json, os, sys, tempfile, time, unittest, zipfile
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from carme.attachments import archive_binary, file_path, save_file, validate_binary
from carme.mcp import MCPError, MCPServer, MCPSession, MCPManager, check_grant
from carme.skills import SkillError, SkillManager
from carme.store import Store
from carme.tools.base import ToolRegistry
from carme.security import digest

class M4(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(prefix='carme-m4-unit-');self.root=Path(self.tmp.name)
        self.env=patch.dict(os.environ,{'CARME_ARTIFACTS_DIR':str(self.root/'cas')});self.env.start()
        self.store=Store(self.root/'state.db');self.skills=SkillManager(self.root/'skills',self.root/'skills.yaml')
        self.cid=self.store.create_conversation(['a','b'])['id']
    def tearDown(self):self.store.close();self.env.stop();self.tmp.cleanup()
    def task(self,bot='a',**kwargs):
        return self.store.create_task(bot,'synthetic goal',conversation_id=self.cid,**kwargs)
    def test_shared_constraint_survives_private_preview(self):
        for i in range(240):self.store.remember('a',f'private{i}',str(i))
        self.store.remember('__shared__','must','SHARED_MUST_SURVIVE')
        text,refs=self.store.memory_context('a')
        self.assertIn('SHARED_MUST_SURVIVE',text);self.assertTrue(any(r['scope']=='user' for r in refs))
    def test_memory_acl_applies_before_query(self):
        self.store.remember('b','secret','PRIVATE_SENTINEL')
        self.assertEqual(self.store.memory_search('a',query='PRIVATE_SENTINEL'),[])
        with self.assertRaisesRegex(ValueError,'denied'):self.store.memory_write('a','user','shared','key','value')
        self.store.memory_grant('a','project','p',write=True)
        self.store.memory_write('a','project','p','key','project secret')
        self.assertEqual(self.store.memory_search('b',query='project secret'),[])
        self.assertEqual(len(self.store.memory_search('a',query='project secret')),1)
    def test_versions_conflict_expiry_revoke_and_task_scope(self):
        ref=self.store.memory_write('a','bot','a','key','v1')
        with self.assertRaisesRegex(ValueError,'conflict'):self.store.memory_write('a','bot','a','key','v2')
        self.store.memory_write('a','bot','a','key','v2',expected_version=1)
        self.assertFalse(self.store.memory_refs_valid([ref]))
        self.store.memory_write('a','bot','a','key','',expected_version=2,revoke=True)
        self.assertFalse(self.store.memory_search('a',key='key'))
        tid=self.task();self.store.memory_write('a','task',tid,'key','task',task_id=tid)
        self.assertFalse(self.store.memory_search('b',task_id=tid))
        self.store.memory_write('a','bot','a','expires','old',expires_at=time.time()+1)
        with patch('carme.store.time.time',return_value=time.time()+2):self.assertFalse(self.store.memory_search('a',key='expires'))
    def test_oversize_constraints_fail_closed(self):
        self.store.memory_write('admin','user','shared','must','x'*13000,constraint=True,admin=True)
        with self.assertRaisesRegex(ValueError,'budget_exceeded'):self.store.memory_context('a')
    def test_skill_install_requires_grant_and_full_paging(self):
        body='---\nname: long\nconstraints: ["Never send without approval"]\n---\n# Intro\n'+'资料'*9000+'\n# Tail\nTAIL_MUST_READ'
        skill=self.skills.install_from_text(body);m=self.skills.snapshot(skill.id)
        with self.assertRaisesRegex(SkillError,'not_granted'):self.skills.page('a',skill.id)
        self.skills.grant('a',skill.id,m['revision']);cursor='';parts=[]
        for _ in range(20):
            page=self.skills.page('a',skill.id,cursor=cursor);parts.append(page['content'])
            self.assertLess(len(json.dumps(page,ensure_ascii=False)),12000);self.assertIn('Never send without approval',page['constraints'])
            cursor=page['next_cursor']
            if not cursor:break
        self.assertTrue(page['complete']);self.assertIn('TAIL_MUST_READ',''.join(parts));self.assertGreater(len(''.join(parts)),12000)
        self.assertEqual(len(self.skills.page('a',skill.id,file='@manifest')['content'])>0,True)
    def test_skill_update_pinned_and_rollback(self):
        skill=self.skills.install_from_text('---\nname: stable\n---\nv1');v1=self.skills.snapshot(skill.id)['revision']
        self.skills.grant('a',skill.id,v1);(skill.path/'SKILL.md').write_text('---\nname: stable\n---\nv2')
        v2=self.skills.snapshot(skill.id)['revision'];self.assertNotEqual(v1,v2)
        self.assertIn('v1',self.skills.page('a',skill.id)['content'])
        self.skills.grant('a',skill.id,v2)
        with self.assertRaisesRegex(SkillError,'stale'):self.skills.page('a',skill.id,cursor=v1+':0')
        self.skills.grant('a',skill.id,v1);self.assertIn('v1',self.skills.page('a',skill.id)['content'])
        with self.assertRaises(SkillError):self.skills.page('b',skill.id)
    def test_manifest_has_all_files_and_rejects_symlinks(self):
        skill=self.skills.install_from_text('---\nname: manifest\n---\nFiles')
        for i in range(230):(skill.path/f'file{i}.txt').write_text(str(i))
        manifest=self.skills.snapshot(skill.id);self.assertEqual(len(manifest['files']),231)
        (skill.path/'link').symlink_to(skill.path/'file1.txt')
        with self.assertRaisesRegex(SkillError,'symlink'):self.skills.snapshot(skill.id)
    def test_skill_immutable_hash_and_path(self):
        skill=self.skills.install_from_text('---\nname: hash\n---\nok');rev=self.skills.snapshot(skill.id)['revision'];self.skills.grant('a',skill.id,rev)
        with self.assertRaises(SkillError):self.skills.page('a',skill.id,file='../secret')
        (self.skills.root/'.versions'/skill.id/rev/'SKILL.md').write_text('tampered')
        with self.assertRaisesRegex(SkillError,'hash_conflict'):self.skills.page('a',skill.id)
    def test_artifact_versions_and_cross_task_acl(self):
        a=self.task();b=self.task('b',meta={'envelope':{}})
        f=archive_binary(self.store,self.cid,'f.txt',b'exact original',task_id=a)
        self.store.update_task_meta(b,{'envelope':{'goal':'review'}})
        with self.assertRaisesRegex(ValueError,'not_granted'):self.store.artifact_access(b,f['id'])
        self.store.artifact_link(b,f['id'],source='delegated');self.assertEqual(self.store.artifact_access(b,f['id'])['sha256'],hashlib.sha256(b'exact original').hexdigest())
        file_path(self.store,f['id']).write_bytes(b'changed')
        with self.assertRaisesRegex(ValueError,'conflict'):self.store.artifact_access(b,f['id'])
    def test_false_completion_is_unverified(self):
        t=self.task();self.store.finish_task(t,'LLM: all tests passed')
        self.assertEqual(self.store.outcome(t)['status'],'unverified')
        with self.assertRaises(SkillError):self.skills.candidate(self.store,t,name='bad',document='bad',private_literals=[])
        with self.assertRaises(ValueError):self.store.accept_outcome(t,digest({}))
    def test_binary_formats_and_independent_content_checks(self):
        from openpyxl import Workbook
        from docx import Document
        from PIL import Image
        from reportlab.pdfgen import canvas
        buffer=io.BytesIO();wb=Workbook();wb.active['A1']=42;wb.save(buffer)
        checks=[{'kind':'xlsx_cell','sheet':'Sheet','cell':'A1','equals':42}]
        self.assertTrue(validate_binary('x.xlsx',buffer.getvalue(),checks)['passed']);checks[0]['equals']=43
        self.assertFalse(validate_binary('x.xlsx',buffer.getvalue(),checks)['passed'])
        buffer=io.BytesIO();doc=Document();doc.add_paragraph('DOCX_SENTINEL');doc.save(buffer)
        self.assertTrue(validate_binary('x.docx',buffer.getvalue(),[{'kind':'text_contains','text':'DOCX_SENTINEL'}])['passed'])
        buffer=io.BytesIO();c=canvas.Canvas(buffer);c.drawString(10,800,'PDF_SENTINEL');c.save()
        self.assertTrue(validate_binary('x.pdf',buffer.getvalue(),[{'kind':'text_contains','text':'PDF_SENTINEL'}])['passed'])
        buffer=io.BytesIO();Image.new('RGB',(16,16),'blue').save(buffer,'PNG')
        self.assertTrue(validate_binary('x.png',buffer.getvalue())['format_valid'])
        with self.assertRaises(ValueError):validate_binary('fake.pdf',b'not pdf')
    def test_archive_traversal_bomb_nested_and_valid_zip(self):
        for name,data in [('../escape',b'x'),('x.zip',b'nested'),('x.txt',b'x'*1000000)]:
            buffer=io.BytesIO()
            with zipfile.ZipFile(buffer,'w',compression=zipfile.ZIP_DEFLATED) as z:z.writestr(name,data)
            with self.assertRaises(ValueError):validate_binary('bad.zip',buffer.getvalue())
        buffer=io.BytesIO()
        with zipfile.ZipFile(buffer,'w') as z:z.writestr('report.txt',b'good')
        check={'kind':'zip_member','name':'report.txt','sha256':hashlib.sha256(b'good').hexdigest()}
        self.assertTrue(validate_binary('good.zip',buffer.getvalue(),[check])['passed'])
    def test_durable_receipt_and_reconciliation(self):
        t=self.task();first=self.store.operation_begin(t,'send','fixed')
        self.store.checkpoint(t,{'messages':[],'next_step':1})
        self.store.close();self.store=Store(self.root/'state.db')
        self.assertEqual(self.store.operation_begin(t,'send','fixed')['status'],'pending')
        self.store.operation_reconcile(t,first['id'],effect='confirmed',receipt={'external_id':'synthetic-42'})
        receipt=self.store.operation_begin(t,'send','fixed');self.assertEqual(receipt['status'],'finished');self.assertIn('synthetic-42',receipt['result'])
        self.assertEqual(self.store.checkpoint(t)['next_step'],1)
    def test_checkpoint_integrity_and_root_budget(self):
        root=self.task(meta={'envelope':{'budget':{'tool_calls':1}}});child=self.task(parent_id=root,meta={'envelope':{'budget':{'tool_calls':1}}})
        self.store.claim_tool_budget(child)
        with self.assertRaisesRegex(ValueError,'budget'):self.store.claim_tool_budget(root)
        self.store.checkpoint(root,{'memory_refs':[]})
        self.store._write('UPDATE task_checkpoints SET payload=? WHERE task_id=?',('{}',root))
        with self.assertRaisesRegex(ValueError,'hash_conflict'):self.store.checkpoint(root)
    def test_mcp_pagination_over_200_and_loop_denied(self):
        session=MCPSession(MCPServer('fixture'));seen=[]
        class Transport:
            async def request(self,method,params,timeout):
                cursor=params.get('cursor');seen.append(cursor)
                if cursor is None:return {'tools':[{'name':f't{i}','inputSchema':{'type':'object'}} for i in range(200)],'nextCursor':'next'}
                return {'tools':[{'name':'tail','inputSchema':{'type':'object'}}]}
        session.transport=Transport();tools=asyncio.run(session.list_tools())
        self.assertEqual(len(tools),201);self.assertEqual(seen,[None,'next'])
        class Loop:
            async def request(self,*a):return {'tools':[],'nextCursor':'same'}
        session.transport=Loop()
        with self.assertRaisesRegex(MCPError,'loop'):asyncio.run(session.list_tools())
    def test_mcp_identity_resource_schema_and_new_tool_denied(self):
        manager=MCPManager(self.root/'mcp.yaml',ToolRegistry());server=MCPServer('fixture',executor='action',command='python')
        manager._servers['fixture']=server
        spec={'name':'read','description':'fixture','inputSchema':{'type':'object','properties':{'path':{'type':'string'}},'required':['path']}}
        manager.catalogs['fixture']=[spec]
        with self.assertRaisesRegex(MCPError,'resource'):manager.grant('a','fixture','read',argument_allowlist={})
        manager.grant('a','fixture','read',argument_allowlist={'path':['/inputs/granted.txt']})
        manager.authorize('a','fixture','read',{'path':'/inputs/granted.txt'})
        for bot,tool,path in [('b','read','/inputs/granted.txt'),('a','new','/inputs/granted.txt'),('a','read','/run/secrets/control')]:
            with self.assertRaises(MCPError):manager.authorize(bot,'fixture',tool,{'path':path})
        spec['inputSchema']['properties']['other']={'type':'string'}
        with self.assertRaisesRegex(MCPError,'changed'):manager.authorize('a','fixture','read',{'path':'/inputs/granted.txt'})
    def test_candidate_sanitization_and_self_eval_rejected(self):
        tid=self.task(meta={'envelope':{'input_artifacts':[]}});self.store.finish_task(tid,'synthetic verified fixture')
        self.store.set_outcome(tid,'verified',{'fixture':True});self.store.accept_outcome(tid,digest({'fixture':True}))
        candidate=self.skills.candidate(self.store,tid,name='candidate',document='---\nname: candidate\n---\nPrivateName sk-FAKE123456789 /Users/Synthetic/private',private_literals=['PrivateName'])
        body=(self.skills.root/'.versions'/candidate['skill_id']/candidate['revision']/'SKILL.md').read_text()
        for secret in ('PrivateName','sk-FAKE123456789','/Users/Synthetic/private'):self.assertNotIn(secret,body)
        with self.assertRaises(SkillError):self.skills.grant('a',candidate['skill_id'],candidate['revision'])
        with self.assertRaises(SkillError):self.skills.publish_candidate(self.store,candidate['id'],test_task_id=tid,bot_ids=['a'],revision=candidate['revision'],privacy_reviewed=True)

    def test_nested_mcp_resource_and_removed_tool_revoke(self):
        manager=MCPManager(self.root/'mcp.yaml',ToolRegistry());manager._servers['s']=MCPServer('s',executor='action',command='python')
        schema={'type':'object','properties':{'options':{'type':'object','properties':{'path':{'type':'string'}},'required':['path']}}}
        manager.catalogs['s']=[{'name':'read','inputSchema':schema}]
        with self.assertRaisesRegex(MCPError,'resource'):manager.grant('a','s','read',argument_allowlist={})
        manager.grant('a','s','read',argument_allowlist={'options':[{'path':'/inputs/only'}]})
        manager.authorize('a','s','read',{'options':{'path':'/inputs/only'}})
        with self.assertRaisesRegex(MCPError,'resource'):manager.authorize('a','s','read',{'options':{'path':'/run/secrets/key'}})
        schema['additionalProperties']=True
        with self.assertRaisesRegex(MCPError,'unsupported'):manager.grant('a','s','read',argument_allowlist={'options':[{'path':'/inputs/only'}]})
        manager.catalogs['s']=[];manager.grant('a','s','read',argument_allowlist={},revoke=True)
        self.assertFalse(manager.grants['a']['s'])

    def test_removed_skill_revokes_and_expired_legacy_recall(self):
        skill=self.skills.install_from_text('---\nname: removed\n---\nbody');rev=self.skills.snapshot(skill.id)['revision']
        self.skills.grant('a',skill.id,rev);self.skills.remove(skill.id)
        self.assertFalse(self.skills.settings()['grants']['a'])
        self.store.memory_write('a','bot','a','expires','fact',expires_at=time.time()+1)
        with patch('carme.store.time.time',return_value=time.time()+2):self.assertFalse(self.store.recall('a','expires'))

    def test_recalled_memory_tracks_revocation_in_task(self):
        from types import SimpleNamespace
        from carme.tools.base import ToolContext
        from carme.tools.memory import RecallTool
        tid=self.task();self.store.remember('a','key','value')
        ctx=ToolContext(agent=SimpleNamespace(id='a'),task_id=tid,store=self.store)
        asyncio.run(RecallTool().run(ctx,key='key'))
        refs=json.loads(self.store.get_task(tid)['meta'])['memory_refs'];self.assertEqual(len(refs),1)
        self.store.forget('a','key');self.assertFalse(self.store.memory_refs_valid(refs))

    def test_bot_proposal_separate_grant_and_source_owner(self):
        from types import SimpleNamespace
        from carme.tools.base import ToolContext,build_registry
        from carme.tools.skill import ProposeSkillTool
        tid=self.task('b');ctx=ToolContext(agent=SimpleNamespace(id='a'),task_id=self.task(),store=self.store)
        self.assertNotIn('propose_skill',build_registry(self.skills).expand(['skill']))
        with self.assertRaisesRegex(SkillError,'source_bot_denied'):
            asyncio.run(ProposeSkillTool(self.skills).run(ctx,tid,name='candidate',document='body',private_literals=[]))

    def test_owned_artifact_and_acceptance_reject_changed_bytes(self):
        tid=self.task(meta={'envelope':{'goal':'exact bytes'}})
        artifact=archive_binary(self.store,self.cid,'original.txt',b'original',task_id=tid)
        report={'artifacts':[{'id':artifact['id'],'sha256':artifact['sha256']}]}
        self.store.set_outcome(tid,'verified',report)
        file_path(self.store,artifact['id']).write_bytes(b'changed')
        with self.assertRaisesRegex(ValueError,'version_conflict'):self.store.artifact_access(tid,artifact['id'])
        with self.assertRaisesRegex(ValueError,'version_conflict'):self.store.accept_outcome(tid,digest(report))

    def test_interrupted_delegation_does_not_create_second_child(self):
        from types import SimpleNamespace
        from carme.tools.base import Tool,ToolContext
        children=[];tid=self.task()
        class DelegatedEffect(Tool):
            name='delegate'
            async def run(inner,ctx,**kwargs):
                children.append('child-with-external-effect')
                raise asyncio.CancelledError()
        registry=ToolRegistry();registry.register(DelegatedEffect())
        def context():
            ctx=ToolContext(agent=SimpleNamespace(id='a',tools=['delegate']),task_id=tid,store=self.store)
            ctx.extras['policy']={'tools':['delegate'],'max_tool_calls':10,'max_output_bytes':65536,'permission_version':'fixture'}
            return ctx
        with self.assertRaises(asyncio.CancelledError):asyncio.run(registry.execute(context(),'delegate',{'agent':'b','goal':'synthetic'}))
        row=self.store._query_one('SELECT * FROM task_operations WHERE task_id=?',(tid,))
        self.assertEqual(row['status'],'pending')
        denied=asyncio.run(registry.execute(context(),'delegate',{'agent':'b','goal':'synthetic'}))
        self.assertIn('reconciliation_required',denied);self.assertEqual(len(children),1)
        self.store.operation_reconcile(tid,row['id'],effect='confirmed',receipt={'child_id':children[0],'result':'confirmed from receipt'})
        resumed=asyncio.run(registry.execute(context(),'delegate',{'agent':'b','goal':'synthetic'}))
        self.assertIn(children[0],resumed);self.assertEqual(len(children),1)


class M4Runtime(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        import test_isolation_m1 as fixture
        self.fixture=fixture.M1();await self.fixture.asyncSetUp()
        self.runtime=self.fixture.runtime;self.store=self.fixture.store
        self.runtime._schedule=lambda *a,**k:None
    async def asyncTearDown(self):await self.fixture.asyncTearDown()
    async def test_frozen_pending_tool_finishes_before_model_continues(self):
        from carme.agents.base import Agent
        from carme.llm import LLMResponse,ToolCall
        from carme.tools.base import Tool
        seen=[]
        class Frozen(Tool):
            name='frozen'
            async def run(inner,ctx,value):seen.append(value);return 'EXACT_RECEIPT'
        self.runtime.registry.register(Frozen());spec=self.fixture.config.agents.get('bot');spec.tools=['frozen']
        tid=await self.runtime.submit('bot','goal');meta=self.runtime._meta(self.store.get_task(tid))
        class Gateway:
            async def chat(inner,messages,**kwargs):
                self.assertEqual(seen,['ORIGINAL_ARGUMENTS']);self.assertEqual(messages[-1]['role'],'tool')
                self.assertEqual(messages[-1]['content'],'EXACT_RECEIPT');return LLMResponse(text='done')
        agent=Agent(spec,self.fixture.config,Gateway(),self.runtime.registry,self.store,self.runtime.skills)
        call=ToolCall('fixed','frozen',{'value':'ORIGINAL_ARGUMENTS'})
        self.store.checkpoint(tid,{'messages':[{'role':'user','content':'goal'},agent._assistant_message(LLMResponse(tool_calls=[call]))],
            'next_step':1,'memory_refs':[],'pending_tool_calls':[{'id':call.id,'name':call.name,'arguments':call.arguments}]})
        self.store.update_task_meta(tid,{'resume_checkpoint':True})
        result=await agent.run('goal',task_id=tid,policy=meta['policy'])
        self.assertEqual(result.status,'done');self.assertFalse(self.store.checkpoint(tid)['pending_tool_calls'])

    async def test_cli_bridge_persists_exact_intent_before_effect(self):
        from carme.agents.base import Agent
        from carme.llm import LLMResponse
        from carme.tools.base import Tool
        seen=[]
        class Effect(Tool):
            name='delegate'
            async def run(inner,ctx,**kwargs):seen.append(kwargs);raise asyncio.CancelledError()
        self.runtime.registry.register(Effect());spec=self.fixture.config.agents.get('bot');spec.engine='pi';spec.tools=['delegate']
        tid=await self.runtime.submit('bot','goal');meta=self.runtime._meta(self.store.get_task(tid))
        class InterruptedGateway:
            async def chat(inner,messages,**kwargs):return await kwargs['cli_tool_execute']('delegate',{'agent':'child','goal':'ORIGINAL'})
        agent=Agent(spec,self.fixture.config,InterruptedGateway(),self.runtime.registry,self.store,self.runtime.skills)
        with self.assertRaises(asyncio.CancelledError):await agent.run('goal',task_id=tid,policy=meta['policy'])
        pending=self.store.checkpoint(tid)['pending_tool_calls'];self.assertEqual(pending[0]['arguments']['goal'],'ORIGINAL')
        operation=self.store._query_one('SELECT * FROM task_operations WHERE task_id=?',(tid,))
        self.store.operation_reconcile(tid,operation['id'],effect='confirmed',receipt={'child_id':'only-child'})
        self.store.update_task_meta(tid,{'resume_checkpoint':True})
        class ResumedGateway:
            async def chat(inner,messages,**kwargs):
                self.assertEqual(len(seen),1);self.assertEqual(messages[-1]['role'],'tool');self.assertIn('only-child',messages[-1]['content'])
                return LLMResponse(text='done')
        agent=Agent(spec,self.fixture.config,ResumedGateway(),self.runtime.registry,self.store,self.runtime.skills)
        result=await agent.run('goal',task_id=tid,policy=meta['policy']);self.assertEqual(result.status,'done');self.assertEqual(len(seen),1)
    async def test_steering_retry_uses_own_contract_hash(self):
        cid=self.store.create_conversation(['bot'])['id'];contract={'expected_outputs':['report.txt']}
        root=await self.runtime.submit_message(cid,'goal','root',envelope=contract)
        self.runtime._active_conversations[cid]=root['task_id'];self.runtime._accepting_input.add(root['task_id'])
        steer=await self.runtime.submit_message(cid,'more input','steer')
        retry=await self.runtime.submit_message(cid,'more input','steer')
        self.assertTrue(steer['steering']);self.assertFalse(retry['created']);self.assertEqual(root['task_id'],retry['task_id'])
        with self.assertRaisesRegex(ValueError,'conflict'):await self.runtime.submit_message(cid,'goal','root',envelope={})
    async def test_bad_envelope_and_child_artifact_access(self):
        cid=self.store.create_conversation(['bot'])['id']
        for payload in ({'input_artifact_ids':[{}]},{'budget':[]},{'input_artifact_ids':['foreign']}):
            with self.assertRaises(ValueError):await self.runtime.submit_message(cid,'bad',str(payload),envelope=payload)
        good=await self.runtime.submit_message(cid,'valid','valid');parent=self.store.get_task(good['task_id'])
        other=self.store.create_task('bot','other',conversation_id=cid)
        artifact=archive_binary(self.store,cid,'other.txt',b'private other task',task_id=other)
        _,meta=self.runtime._task_agent_snapshot('bot',parent)
        child=self.store.create_task('bot','child',parent_id=parent['id'],conversation_id=cid,meta=meta)
        with self.assertRaisesRegex(ValueError,'artifact_version_not_granted'):self.runtime.bind_envelope(child,{'input_artifact_ids':[artifact['id']]})
    async def test_resume_denies_descendant_effect_and_changed_policy(self):
        cid=self.store.create_conversation(['bot'])['id'];result=await self.runtime.submit_message(cid,'goal','root')
        tid=result['task_id'];self.store.finish_task(tid,'interrupted',status='failed')
        self.store.checkpoint(tid,{'messages':[],'next_step':1,'memory_refs':[]})
        child=self.store.create_task('bot','child',parent_id=tid,conversation_id=cid)
        operation=self.store.operation_begin(child,'mcp__synthetic__send','digest')
        with self.assertRaisesRegex(ValueError,'reconciliation'):await self.runtime.resume(tid)
        self.store.operation_reconcile(child,operation['id'],effect='confirmed',receipt={'synthetic_external_id':'one'})
        self.store.memory_grant('bot','project','new',read=True)
        with self.assertRaisesRegex(ValueError,'permission_version_changed'):await self.runtime.resume(tid)

if __name__=='__main__':unittest.main(verbosity=2)
